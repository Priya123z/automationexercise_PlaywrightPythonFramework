from openpyxl import Workbook, load_workbook
import os
import json
import csv

class DataWriter:

    @staticmethod
    def write_xlsx(file_path, data):

        # Check if file exists
        if os.path.exists(file_path):

            workbook = load_workbook(file_path)
            sheet = workbook.active

        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "TestData"

            # Write header only once
            sheet.append(list(data.keys()))

        # Write row values
        sheet.append(list(data.values()))

        # Save file
        workbook.save(file_path)

        print(f"Data written to {file_path}")

    @staticmethod
    def write_csv(file_path, data):

        file_exists = os.path.exists(file_path)

        with open(file_path, mode="a", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(file, fieldnames=data.keys())

            # Write header only once
            if not file_exists:
                writer.writeheader()

            writer.writerow(data)

        print("Data written to CSV")

    @staticmethod
    def write_json(file_path, data):

        existing = []

        if os.path.exists(file_path):
            with open(file_path, "r") as file:
                try:
                    existing = json.load(file)
                except:
                    existing = []

        existing.append(data)

        with open(file_path, "w") as file:
            json.dump(existing, file, indent=4)

        print("Data written to JSON")

